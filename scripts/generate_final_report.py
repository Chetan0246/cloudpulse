#!/usr/bin/env python3
"""
CloudPulse v2 — Final Review Technical Report Generator.
Generates a comprehensive ~30-page academic and engineering report matching
the formatting and rubric of the VIT BCSE355L Course Project template.
"""

import os
import subprocess
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

NAVY = "1E3A8A"
BLUE = "2563EB"
DARK_GRAY = "1F2937"
LIGHT_GRAY = "F8FAFC"
BORDER_GRAY = "CBD5E1"
CODE_BG = "F1F5F9"
CALLOUT_BG = "EFF6FF"

def set_cell_background(cell, hex_color):
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    cell._tc.get_or_add_tcPr().append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def set_table_borders(table, color="CBD5E1", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:left w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:right w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideV w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def add_header_block(doc):
    p0 = doc.add_paragraph()
    p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r0 = p0.add_run("School of Computer Science Engineering and Information Systems\n")
    r0.font.name = "Tahoma"
    r0.font.size = Pt(13)
    r0.font.bold = True
    r0.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)

    r1 = p0.add_run("B.Tech  Information Technology\n")
    r1.font.name = "Tahoma"
    r1.font.size = Pt(12)
    r1.font.bold = True

    r2 = p0.add_run("BCSE355L : Cloud Architecture Design\n")
    r2.font.name = "Tahoma"
    r2.font.size = Pt(12)
    r2.font.bold = True

    r3 = p0.add_run("FALL 2026-2027\n")
    r3.font.name = "Tahoma"
    r3.font.size = Pt(11)
    r3.font.bold = True

    r4 = p0.add_run("Course Project — Final Review Report (Review III / Final Evaluation)\n")
    r4.font.name = "Tahoma"
    r4.font.size = Pt(12)
    r4.font.bold = True
    r4.font.color.rgb = RGBColor(0x25, 0x63, 0xEB)

    r5 = p0.add_run("(Review Date : October 2026)\n")
    r5.font.name = "Tahoma"
    r5.font.size = Pt(10)
    r5.font.italic = True

def add_metadata_box(doc):
    tbl = doc.add_table(rows=3, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl, color="1E3A8A", sz="8")
    
    # Row 0: Team Name
    c0 = tbl.cell(0, 0)
    set_cell_background(c0, "F8FAFC")
    set_cell_margins(c0, top=120, bottom=120, left=180, right=180)
    p0 = c0.paragraphs[0]
    r0a = p0.add_run("TEAM NAME : ")
    r0a.font.name = "Tahoma"
    r0a.font.bold = True
    r0a.font.size = Pt(11)
    r0b = p0.add_run("Nebula Knight")
    r0b.font.name = "Tahoma"
    r0b.font.bold = True
    r0b.font.size = Pt(11)
    r0b.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)

    # Row 1: Member details
    c1 = tbl.cell(1, 0)
    set_cell_background(c1, "FFFFFF")
    set_cell_margins(c1, top=120, bottom=120, left=180, right=180)
    p1 = c1.paragraphs[0]
    r1a = p1.add_run("Team Member with Reg # and Name : ")
    r1a.font.name = "Tahoma"
    r1a.font.bold = True
    r1a.font.size = Pt(10.5)
    r1b = p1.add_run("Moorthy Chetan; 24BIT0246; Mobile: 9441522812; Email: moorthychetan06@gmail.com")
    r1b.font.name = "Tahoma"
    r1b.font.size = Pt(10.5)

    # Row 2: Title
    c2 = tbl.cell(2, 0)
    set_cell_background(c2, "F8FAFC")
    set_cell_margins(c2, top=140, bottom=140, left=180, right=180)
    p2 = c2.paragraphs[0]
    r2a = p2.add_run("Project Title : ")
    r2a.font.name = "Tahoma"
    r2a.font.bold = True
    r2a.font.size = Pt(11)
    r2b = p2.add_run("CloudPulse v2: Autonomous Cloud Reliability and Self-Healing Control Plane Using Real AWS Multi-Service Architecture")
    r2b.font.name = "Tahoma"
    r2b.font.bold = True
    r2b.font.size = Pt(11)
    r2b.font.color.rgb = RGBColor(0x25, 0x63, 0xEB)

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

def add_h1(doc, title, page_break=False):
    if page_break:
        doc.add_page_break()
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(title)
    r.font.name = "Tahoma"
    r.font.size = Pt(14)
    r.font.bold = True
    r.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)
    return p

def add_h2(doc, title):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(11)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(title)
    r.font.name = "Tahoma"
    r.font.size = Pt(12)
    r.font.bold = True
    r.font.color.rgb = RGBColor(0x25, 0x63, 0xEB)
    return p

def add_h3(doc, title):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(title)
    r.font.name = "Tahoma"
    r.font.size = Pt(10.5)
    r.font.bold = True
    r.font.color.rgb = RGBColor(0x33, 0x41, 0x55)
    return p

def add_body_p(doc, text, bold_prefix=""):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        rb = p.add_run(bold_prefix)
        rb.font.name = "Tahoma"
        rb.font.size = Pt(10)
        rb.font.bold = True
        rb.font.color.rgb = RGBColor(0x1F, 0x29, 0x37)
    r = p.add_run(text)
    r.font.name = "Tahoma"
    r.font.size = Pt(10)
    r.font.color.rgb = RGBColor(0x1F, 0x29, 0x37)
    return p

def add_bullet(doc, text, bold_prefix=""):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        rb = p.add_run(bold_prefix)
        rb.font.name = "Tahoma"
        rb.font.size = Pt(9.5)
        rb.font.bold = True
        rb.font.color.rgb = RGBColor(0x1F, 0x29, 0x37)
    r = p.add_run(text)
    r.font.name = "Tahoma"
    r.font.size = Pt(9.5)
    r.font.color.rgb = RGBColor(0x1F, 0x29, 0x37)
    return p

def add_callout(doc, text, title="OPERATIONAL ARCHITECTURE NOTE"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    c = tbl.cell(0, 0)
    set_cell_background(c, CALLOUT_BG)
    set_cell_margins(c, top=100, bottom=100, left=150, right=150)
    
    # Left border only (blue)
    tcPr = c._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:top w:val="none"/>'
        f'<w:left w:val="single" w:sz="24" w:space="0" w:color="{BLUE}"/>'
        f'<w:bottom w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)
    
    p = c.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    rt = p.add_run(f"[{title}] ")
    rt.font.name = "Tahoma"
    rt.font.size = Pt(9.5)
    rt.font.bold = True
    rt.font.color.rgb = RGBColor(0x25, 0x63, 0xEB)
    
    r = p.add_run(text)
    r.font.name = "Tahoma"
    r.font.size = Pt(9.5)
    r.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def add_code(doc, code_str, caption=""):
    if caption:
        cp = doc.add_paragraph()
        cp.paragraph_format.space_before = Pt(6)
        cp.paragraph_format.space_after = Pt(2)
        rc = cp.add_run(f"Code Listing: {caption}")
        rc.font.name = "Tahoma"
        rc.font.size = Pt(9)
        rc.font.bold = True
        rc.font.color.rgb = RGBColor(0x47, 0x55, 0x69)
        
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    c = tbl.cell(0, 0)
    set_cell_background(c, CODE_BG)
    set_cell_margins(c, top=90, bottom=90, left=120, right=120)
    set_table_borders(tbl, color="CBD5E1", sz="4")
    
    p = c.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.05
    r = p.add_run(code_str.strip())
    r.font.name = "Consolas"
    r.font.size = Pt(8.5)
    r.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def add_styled_table(doc, headers, data, col_widths=None):
    tbl = doc.add_table(rows=len(data) + 1, cols=len(headers))
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl, color="CBD5E1", sz="4")
    
    # Header Row
    hdr_cells = tbl.rows[0].cells
    for i, h in enumerate(headers):
        hdr_cells[i].text = h
        set_cell_background(hdr_cells[i], NAVY)
        set_cell_margins(hdr_cells[i], top=100, bottom=100, left=120, right=120)
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in p.runs:
            run.font.name = "Tahoma"
            run.font.size = Pt(9.5)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            
    # Data Rows
    for r_idx, row in enumerate(data):
        row_cells = tbl.rows[r_idx + 1].cells
        bg_col = LIGHT_GRAY if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate(row):
            row_cells[c_idx].text = str(val)
            set_cell_background(row_cells[c_idx], bg_col)
            set_cell_margins(row_cells[c_idx], top=80, bottom=80, left=100, right=100)
            p = row_cells[c_idx].paragraphs[0]
            if c_idx == 0:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            for run in p.runs:
                run.font.name = "Tahoma"
                run.font.size = Pt(9)
                run.font.color.rgb = RGBColor(0x1F, 0x29, 0x37)

    if col_widths:
        for row in tbl.rows:
            for idx, width in enumerate(col_widths):
                row.cells[idx].width = Inches(width)

    doc.add_paragraph().paragraph_format.space_after = Pt(6)
    return tbl

print("Helper definitions loaded successfully.")
